"""CampusScraper tests on the fictional Sunrise Institute of Technology site (tests_campus/site)."""
import json
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from campusscraper import web
from campusscraper.cli import main
from campusscraper.config import ScrapeConfig
from campusscraper.crawl import Crawler
from campusscraper.extract import extract_page
from campusscraper.models import Page
from campusscraper.sources import read_sources, write_template

CSE = {"Dr. Anitha Rao", "Prof. K. Srinivas", "Ms. Priya Sharma", "Mr. Rajesh Kumar", "Dr. Farah Siddiqui", "Mrs. Lakshmi Devi"}
MECH = {"Venkata Ramana Murthy", "Sneha Kulkarni"}
NON_TEACHING = {"Mr. Ramesh Babu", "Mr. Suresh Naik"}
ALUMNI = {"Dr. Sanjay Patil", "Mr. Nikhil Deshmukh", "Dr. Meera Joshi", "Ms. Pooja Kale", "Mr. Rohan Mehta",
          "Dr. Vivek Kulkarni", "Dr. Asha Bhide", "Mr. Tushar Gawande"}
STUDENT_NAMES = ("Rahul Verma", "Sravani Reddy", "Kiran Teja", "Naveen Kumar", "Divya Menon", "Arjun Patel")
STUDENT_CONTACTS = ("90000 33333", "90000 44444", "98765 00001", "90000 55555", "divya.m22", "rahul.v21", "sravani.r21",
                    "70000 66666")


def cfg(**kw):
    base = dict(max_pages=60, delay_seconds=0, allow_private_hosts=True, render="never", retries=0)
    base.update(kw)
    return ScrapeConfig(**base)


def by_name(result):
    return {d.get("name"): d for d in result.doctors}


def test_crawl_finds_all_teaching_faculty(site_url, hits):
    result = Crawler(cfg()).run([site_url])
    fac = {d.get("name") for d in result.faculty}
    assert CSE | MECH <= fac, (CSE | MECH) - fac
    assert not fac & NON_TEACHING
    people = by_name(result)
    assert NON_TEACHING <= set(people)  # kept in 'All people'
    names = [d.get("name") for d in result.doctors]
    assert len(names) == len(set(names))
    assert not any(h.startswith("/admin") for h in hits)  # robots.txt
    assert "Dr. Secret Person" not in people


def test_profile_fields(site_url):
    docs = by_name(Crawler(cfg()).run([site_url]))
    ar = docs["Dr. Anitha Rao"]
    assert ar.page_type == "profile"
    assert ar.get("designation") == "Professor & Head"
    assert ar.get("department") == "Computer Science and Engineering"
    assert ar.get("qualifications")[:2] == ["PhD (Computer Science)", "MTech (CSE)"]
    assert ar.get("experience_years") == "22"
    assert ar.get("email") == "anitha.rao@sunrise-it.ac.in"
    assert ar.get("phone") == "+91 98480 11111"
    assert "hodcse@sunrise-it.ac.in" in (ar.get("contact_email") or "")  # department mailbox, not hers
    assert ar.get("research_areas") == ["Machine Learning", "Natural Language Processing", "Information Retrieval"]
    assert ar.get("subjects") == ["Data Structures", "Compiler Design"]
    assert "Senior Member, IEEE" in ar.get("memberships")
    assert ar.get("google_scholar", "").startswith("https://scholar.google.co.in/citations?user=")
    assert ar.get("orcid") == "0000-0002-1825-0097"
    assert ar.get("vidwan") == "https://vidwan.inflibnet.ac.in/profile/111111"
    assert ar.get("institution") == "Sunrise Institute of Technology, Hyderabad" and ar.get("city") == "Hyderabad"
    # header/footer switchboard never becomes a person's contact
    assert "2345 6789" not in json.dumps(ar.to_dict())


def test_other_layouts(site_url):
    docs = by_name(Crawler(cfg()).run([site_url]))
    ks = docs["Prof. K. Srinivas"]  # JSON-LD Person with sameAs
    assert ks.get("orcid") == "0000-0001-5109-3700"
    assert ks.get("linkedin") == "https://www.linkedin.com/in/ksrinivas-example"
    assert ks.get("research_areas")[:2] == ["Cloud Computing", "Distributed Systems"]
    assert ks.get("department") == "Computer Science and Engineering"
    ps = docs["Ms. Priya Sharma"]  # dt/dd
    assert ps.get("scopus_id") == "57212345678" and ps.get("experience_years") == "6"
    assert ps.get("research_areas") == ["Deep Learning", "Computer Vision"]
    fs = docs["Dr. Farah Siddiqui"]
    assert fs.get("designation") == "Associate Professor (Sr. Grade)"
    vm = docs["Venkata Ramana Murthy"]  # table row without honorific
    assert vm.get("department") == "Mechanical Engineering" and vm.get("email") == "vr.murthy@sunrise-it.ac.in"
    assert docs["Mr. Rajesh Kumar"].get("department") == "Computer Science and Engineering"  # 'CSE' -> canonical
    assert docs["Mrs. Lakshmi Devi"].get("department") == "Computer Science and Engineering"  # from page heading


def test_department_filter(site_url):
    result = Crawler(cfg(departments=["CSE"])).run([site_url])
    fac = {d.get("name") for d in result.faculty}
    assert fac == CSE
    assert MECH <= {d.get("name") for d in result.doctors}


def test_student_offices_and_no_student_contacts(site_url):
    result = Crawler(cfg()).run([site_url])
    offices = {o.office_type: o for o in result.offices}
    assert set(offices) == {"placement", "higher_studies", "student_chapter", "alumni"}
    tp = offices["placement"]
    assert tp.emails == ["tpo@sunrise-it.ac.in"]
    assert "+91 40 2345 6800" in tp.phones
    assert "Dr. Mohan Krishna, Training and Placement Officer" in tp.coordinators
    assert any("42 students admitted" in l for l in tp.higher_studies_lines)
    assert tp.skipped_student_lines >= 2
    hs = offices["higher_studies"]
    assert hs.emails == ["higherstudies@sunrise-it.ac.in"] and hs.coordinators[0].startswith("Dr. Farah Siddiqui")
    ieee = offices["student_chapter"]
    assert ieee.emails == ["ieee@sunrise-it.ac.in"] and ieee.phones == []
    assert ieee.coordinators == ["Prof. K. Srinivas, Faculty Advisor"]
    assert offices["alumni"].coordinators == ["Mrs. Geetha Rani, Alumni Relations Officer"]
    # Nothing about individual students anywhere in the output.
    dump = json.dumps([o.to_full_dict() for o in result.offices] + [d.to_full_dict() for d in result.doctors])
    for s in STUDENT_NAMES + STUDENT_CONTACTS:
        assert s not in dump, s


def test_cli_scrape_writes_excel(site_url, tmp_path, capsys):
    out = tmp_path / "college.xlsx"
    rc = main(["scrape", site_url, "--out", str(out), "--delay", "0", "--render", "never", "--max-pages", "60",
               "--allow-local", "--departments", "CSE"])
    assert rc == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Faculty", "All people", "Student offices", "Evidence", "Pages"]
    head = [c.value for c in wb["Faculty"][1]]
    assert head[:4] == ["Name", "Designation", "Department", "Specialization"] and "Google Scholar" in head
    names = {r[0].value for r in wb["Faculty"].iter_rows(min_row=2)}
    assert names == CSE
    otypes = {r[0].value for r in wb["Student offices"].iter_rows(min_row=2)}
    assert otypes == {"placement", "higher_studies", "student_chapter", "alumni"}
    assert "student offices" in capsys.readouterr().out


def test_cli_page_and_file(tmp_path, capsys):
    p = tmp_path / "p.html"
    p.write_text("<!-- source: https://c.example/faculty/dr-x-y -->\n<html><body><h1>Dr. Xavier Yadav</h1>"
                 "<p>Assistant Professor, Department of Information Technology</p><p>Ph.D, M.Tech (IT)</p></body></html>")
    assert main(["file", str(p)]) == 0
    out = capsys.readouterr().out
    assert "Dr. Xavier Yadav" in out and "Information Technology" in out


def test_template_round_trip_and_aliases(tmp_path):
    t = write_template(tmp_path / "t.xlsx", rows=[{
        "college_name": "Sunrise Institute", "faculty_page_urls": "https://www.sunrise.ac.in/cse/faculty",
        "student_office_urls": "https://www.sunrise.ac.in/training-and-placement",
        "sample_profile_url": "https://www.sunrise.ac.in/faculty/dr-a-rao", "departments": "CSE, IT"}])
    sources, problems = read_sources(t)
    assert not problems and sources[0].institution == "Sunrise Institute"
    assert sources[0].patterns == ["/faculty/"] and sources[0].departments == ["CSE", "IT"]
    assert len(sources[0].urls) == 3
    csv_path = tmp_path / "list.csv"
    csv_path.write_text("college,website,placement_url\nA College,www.a-college.ac.in,\nB College,,b.ac.in/tpo\n")
    sources, _ = read_sources(csv_path)
    assert [s.urls for s in sources] == [["https://www.a-college.ac.in/"], ["https://b.ac.in/tpo"]]


def test_batch_with_resume(site_url, tmp_path, capsys, monkeypatch):
    monkeypatch.setenv("CAMPUSSCRAPER_ALLOW_PRIVATE", "1")
    csv_path = tmp_path / "colleges.csv"
    csv_path.write_text("college_name,faculty_page_urls,student_office_urls,departments\n"
                        f"Sunrise IT,{site_url}/departments/cse/faculty,{site_url}/training-and-placement,CSE\n")
    out = tmp_path / "batch.xlsx"
    args = ["batch", str(csv_path), "--out", str(out), "--delay", "0", "--render", "never", "--pages-per-site", "30"]
    assert main(args) == 0
    wb = load_workbook(out)
    assert "Colleges" in wb.sheetnames
    row = [c.value for c in wb["Colleges"][2]]
    assert row[1] == "Sunrise IT" and row[2] == "done"
    names = {r[0].value for r in wb["Faculty"].iter_rows(min_row=2)}
    assert CSE <= names
    offices = [r for r in wb["Student offices"].iter_rows(min_row=2, values_only=True)]
    assert any(r[0] == "placement" and r[9] == "Sunrise IT" for r in offices)
    capsys.readouterr()
    assert main(args) == 0
    assert "resumed" in capsys.readouterr().out


def test_web_ui_job_flow(site_url, tmp_path, monkeypatch):
    monkeypatch.setattr(web, "OUT_DIR", tmp_path)
    monkeypatch.setenv("CAMPUSSCRAPER_ALLOW_PRIVATE", "1")
    client = TestClient(web.app)
    assert "CampusScraper" in client.get("/").text
    jid = client.post("/api/jobs", json={"urls": [site_url], "delay_seconds": 0, "render": "never", "max_pages": 60,
                                         "departments": "CSE"}).json()["id"]
    for _ in range(150):
        st = client.get(f"/api/jobs/{jid}").json()
        if st["state"] != "running":
            break
        time.sleep(0.1)
    assert st["state"] == "done", st
    assert st["progress"]["faculty"] == len(CSE)
    assert {d["name"] for d in st["doctors"]} == CSE
    assert {o["office_type"] for o in st["offices"]} == {"placement", "higher_studies", "student_chapter", "alumni"}
    r = client.get(f"/api/jobs/{jid}/download.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert "Dr. Anitha Rao" in client.get(f"/api/jobs/{jid}/download.csv").text
    j = json.loads(client.get(f"/api/jobs/{jid}/download.json").text)
    assert j["stats"]["faculty"] == len(CSE) and len(j["offices"]) == 4
    t = client.post("/api/page", json={"url": site_url + "/training-and-placement", "render": "never"}).json()
    assert t["offices"][0]["office_type"] == "placement"
    assert "90000 33333" not in json.dumps(t)  # a student's number never reaches the screen


def test_web_batch_upload(site_url, tmp_path, monkeypatch):
    import base64
    monkeypatch.setattr(web, "OUT_DIR", tmp_path)
    monkeypatch.setenv("CAMPUSSCRAPER_ALLOW_PRIVATE", "1")
    client = TestClient(web.app)
    tpl = client.get("/api/template.xlsx")
    assert tpl.status_code == 200 and tpl.content[:2] == b"PK"
    csv_text = f"college_name,faculty_page_urls\nSunrise IT,{site_url}/departments/mech/faculty\n"
    body = {"filename": "c.csv", "content_base64": base64.b64encode(csv_text.encode()).decode()}
    chk = client.post("/api/sources/check", json=body).json()
    assert chk["usable"] == 1 and chk["sources"][0]["institution"] == "Sunrise IT"
    bid = client.post("/api/batch", json=body | {"delay_seconds": 0, "render": "never", "pages_per_site": 10}).json()["id"]
    for _ in range(150):
        st = client.get(f"/api/batch/{bid}").json()
        if st["state"] != "running":
            break
        time.sleep(0.1)
    assert st["state"] == "done", st
    assert st["sites"][0]["faculty"] >= 2
    docs = client.get(f"/api/batch/{bid}/doctors").json()["doctors"]
    assert MECH <= {d["name"] for d in docs}
    assert client.get(f"/api/batch/{bid}/download.xlsx").content[:2] == b"PK"


def test_student_card_is_not_a_person_record():
    html = ("<html><body><h1>Student Coordinators</h1><div class='c'><h4>Mr. Rohit Sen</h4><p>Student Coordinator, "
            "III Year CSE</p><p>9876512345</p></div><div class='c'><h4>Dr. Asha Iyer</h4><p>Associate Professor, "
            "Faculty Coordinator</p></div></body></html>")
    pr = extract_page(Page("https://c.example/student-activities", "https://c.example/student-activities", 200, html))
    names = {d.get("name") for d in pr.doctors}
    assert "Mr. Rohit Sen" not in names and "Dr. Asha Iyer" in names
    assert "+91 98765 12345" not in json.dumps([o.to_dict() for o in pr.offices])


def test_table_with_name_and_designation_in_one_column():
    html = """<html><body><h2>Faculty Details - Department of ECE</h2><table>
<tr><th>S.No</th><th>Photo</th><th>Name &amp; Designation</th><th>Qualification</th><th>Experience</th><th>E-mail ID</th></tr>
<tr><td>1</td><td><img src=a.jpg></td><td>Dr. M. Suresh<br>Professor &amp; HOD</td><td>Ph.D</td><td>24</td><td>suresh.m@abc.edu.in</td></tr>
<tr><td>2</td><td><img src=b.jpg></td><td>Mr. P. Kiran<br>Assistant Professor</td><td>M.Tech</td><td>8</td><td>kiran.p@abc.edu.in</td></tr>
<tr><td>3</td><td><img src=c.jpg></td><td>Mrs. S. Rekha<br>Assistant Professor</td><td>M.E</td><td>5</td><td>rekha.s@abc.edu.in</td></tr>
</table></body></html>"""
    pr = extract_page(Page("https://abc.edu.in/ece/faculty", "https://abc.edu.in/ece/faculty", 200, html))
    rows = {d.get("name"): d for d in pr.doctors}
    assert set(rows) == {"Dr. M. Suresh", "Mr. P. Kiran", "Mrs. S. Rekha"}
    assert rows["Dr. M. Suresh"].get("designation") == "Professor & HOD"
    assert rows["Mrs. S. Rekha"].get("qualifications") == ["ME"]
    assert all(d.get("department") == "Electronics and Communication Engineering" and d.is_relevant for d in pr.doctors)


def _wait(client, url, done, tries=200):
    for _ in range(tries):
        st = client.get(url).json()
        if done(st):
            return st
        time.sleep(0.1)
    raise AssertionError(st)


def test_portal_history_directories_dashboard_and_resume(site_url, tmp_path, monkeypatch):
    import base64
    monkeypatch.setattr(web, "OUT_DIR", tmp_path)
    monkeypatch.setenv("CAMPUSSCRAPER_ALLOW_PRIVATE", "1")
    monkeypatch.delenv("CAMPUSSCRAPER_DB", raising=False)
    web._STORE = None
    client = TestClient(web.app)
    path = write_template(tmp_path / "p.xlsx", [
        {"college_name": "Sunrise CSE", "faculty_page_urls": f"{site_url}/departments/cse/faculty",
         "student_office_urls": f"{site_url}/training-and-placement"},
        {"college_name": "Sunrise Mech", "faculty_page_urls": f"{site_url}/departments/mech/faculty",
         "student_office_urls": f"{site_url}/higher-studies-cell"}])
    payload = {"filename": "p.xlsx", "content_base64": base64.b64encode(path.read_bytes()).decode(),
               "pages_per_site": 8, "delay_seconds": 0, "render": "never", "use_sitemaps": False}
    bid = client.post("/api/batch", json=payload).json()["id"]
    _wait(client, f"/api/batch/{bid}", lambda s: s["state"] == "done", tries=300)
    runs = client.get("/api/runs").json()["runs"]
    assert runs[0]["id"] == bid and runs[0]["state"] == "done"
    assert runs[0]["faculty_saved"] > 0 and runs[0]["offices_saved"] >= 2
    assert len(client.get(f"/api/runs/{bid}").json()["sites"]) == 2
    d = client.get("/api/directory").json()
    assert d["total"] > 0 and {"Sunrise CSE", "Sunrise Mech"} <= set(d["facets"]["colleges"])
    assert "Mechanical Engineering" in d["facets"]["departments"]
    mech = client.get("/api/directory", params={"college": "Sunrise Mech"}).json()
    assert {x["name"] for x in mech["doctors"]} >= MECH and all(x["source_institution"] == "Sunrise Mech" for x in mech["doctors"])
    cse = client.get("/api/directory", params={"department": "Computer Science and Engineering"}).json()
    assert cse["total"] >= 4 and all(x["department"] == "Computer Science and Engineering" for x in cse["doctors"])
    assert client.get("/api/directory", params={"has_profile": "true"}).json()["total"] >= 1
    assert client.get("/api/directory", params={"q": "zzzz-nothing"}).json()["total"] == 0
    staff = client.get("/api/directory", params={"all_people": "true", "college": "Sunrise Mech"}).json()
    assert NON_TEACHING <= {x["name"] for x in staff["doctors"]}
    offs = client.get("/api/offices").json()
    assert {o["office_type"] for o in offs["offices"]} >= {"placement", "higher_studies"}
    tp = client.get("/api/offices", params={"office_type": "placement"}).json()["offices"]
    assert tp[0]["emails"] == ["tpo@sunrise-it.ac.in"] and tp[0]["college"] == "Sunrise CSE"
    assert client.get("/api/offices", params={"has_figures": "true"}).json()["total"] >= 2
    dash = client.get("/api/dashboard").json()
    assert dash["totals"]["faculty"] == d["total"] and dash["totals"]["offices"] == offs["total"]
    assert dash["totals"]["placement_cells"] >= 1 and len(dash["colleges"]) == 2
    assert client.get("/api/directory/download.xlsx").content[:2] == b"PK"
    assert "Sunrise" in client.get("/api/directory/download.csv").text
    assert client.get("/api/offices/download.xlsx").content[:2] == b"PK"
    assert "tpo@sunrise-it.ac.in" in client.get("/api/offices/download.csv").text
    assert client.get(f"/api/runs/{bid}/download.xlsx").content[:2] == b"PK"
    for s in STUDENT_NAMES + STUDENT_CONTACTS:  # the history holds no students either
        assert s not in json.dumps(offs) + json.dumps(staff)
    # app restart while a batch was going: it shows as interrupted and can be resumed
    web.store().update_run(bid, state="running")
    web.BATCHES.pop(bid)
    web._STORE = None
    assert client.get(f"/api/runs/{bid}").json()["state"] == "interrupted"
    r = client.post(f"/api/runs/{bid}/resume").json()
    assert r["id"] == bid and r["colleges"] == 2
    st = _wait(client, f"/api/batch/{bid}", lambda s: s["state"] == "done", tries=300)
    assert {s["status"] for s in st["sites"]} == {"resumed"}  # finished colleges were not crawled again
    # single crawl: saved, and can be run again
    jid = client.post("/api/jobs", json={"urls": [site_url + "/departments/cse/faculty"], "delay_seconds": 0,
                                         "render": "never", "max_pages": 6, "use_sitemaps": False}).json()["id"]
    _wait(client, f"/api/jobs/{jid}", lambda s: s["state"] == "done")
    again = client.post(f"/api/runs/{jid}/rerun").json()["id"]
    assert again != jid
    _wait(client, f"/api/jobs/{again}", lambda s: s["state"] == "done")
    # delete
    assert client.delete(f"/api/runs/{bid}").json()["ok"]
    assert all(x["id"] != bid for x in client.get("/api/runs").json()["runs"])
    web._STORE = None


def test_alumni_are_never_people(site_url):
    """Home page with 'Prestigious Alumni' and 'Faculty Members' carousels side by side, a testimonial slider,
    and a 'Notable Alumni' page whose alumni are professors elsewhere: only the college's own faculty are kept."""
    result = Crawler(cfg()).run([site_url])
    names = {d.get("name") for d in result.doctors}
    assert not names & ALUMNI, names & ALUMNI
    assert CSE | MECH <= {d.get("name") for d in result.faculty}
    home = extract_page(Page(site_url + "/", site_url + "/", 200,
                             (__import__("pathlib").Path(__file__).parent / "site" / "index.html").read_text()))
    assert {d.get("name") for d in home.doctors} == {"Dr. Anitha Rao", "Prof. K. Srinivas"}


def test_page_headings_are_not_people():
    from campusscraper.profile import plain_person_name
    for text in ("About Us", "Organization Chart", "Principal's Message", "Vision Mission", "Photo Gallery",
                 "Anti Ragging", "Fee Structure", "Governing Body", "Academic Calendar"):
        assert plain_person_name(text) is None, text
    html = ("<html><body><h1>Administration</h1><div class='grid'>"
            "<div class='c'><h4>About Us</h4><p>Principal</p><a href='/about-us'>read</a></div>"
            "<div class='c'><h4>Organization Chart</h4><p>Principal</p><a href='/org-chart'>read</a></div>"
            "<div class='c'><h4>Smt. M. V. Patil</h4><p>Principal</p></div></div></body></html>")
    pr = extract_page(Page("https://c.example/administration", "https://c.example/administration", 200, html))
    assert {d.get("name") for d in pr.doctors} == {"Smt. M. V. Patil"}
