"""Workflow simulations: real crawler and Excel, controlled search/model replies.

These tests do NOT assert a live LLM's semantic accuracy or real-site recall.
"""
import json
import threading
from dataclasses import asdict
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from campusscraper.config import ScrapeConfig
from campusscraper.discovery_services import DiscoveryDB, ModelRouter, ServiceWait
from campusscraper.models import Page, PersonRecord
from campusscraper.prompt_discovery import Plan, Decision, Report, Limits, Refinement
from campusscraper.prompt_jobs import PromptIn, PromptManager, export_sources
from campusscraper.store import Store
from campusscraper.verified_discovery import (
    PlanReview, RecordBatchDecision, _check_page, discover_verified, filter_records,
)
from test_prompt_discovery import Fetch, Search, ScriptedModel, plan, decision, URL, HTML


REVIEW = PlanReview(outcome="ok", questions=[], reason="Requirements preserved")


@pytest.mark.parametrize("prompt", ["Engineering faculty chahiye Mumbai mein, Navi Mumbai nahi",
    "COEP and PICT official people pages", "Universities anywhere with marine biology teams",
    "Placement office contacts at autonomous colleges", "Find chemistry professors with published emails"])
def test_free_form_prompts_reach_planner_and_auditor_intact(prompt):
    model = ScriptedModel(plan(requested_institutions=1), REVIEW, decision())
    report = discover_verified(prompt, model, Search(), Fetch(), Limits(rounds=1))
    assert report.candidates
    assert all(call[2]["original_prompt"] == prompt for call in model.calls)


def test_audit_catches_silently_dropped_exclusion():
    model = ScriptedModel(plan(), PlanReview(outcome="clarify", questions=["Should Navi Mumbai be excluded?"], reason="The plan omitted the exclusion"))
    search = Search()
    report = discover_verified("Mumbai, exclude Navi Mumbai", model, search, Fetch())
    assert report.status == "clarify" and search.queries == []


@pytest.mark.parametrize("scope", ["students", "mixed", "other"])
def test_unsupported_requests_do_not_search(scope):
    model = ScriptedModel(plan(action="unsupported", scope=scope))
    search = Search()
    report = discover_verified("Find individual student emails", model, search, Fetch())
    assert report.status == "unsupported" and not search.queries


def test_review_can_reject_student_plan_even_when_planner_wrong():
    model = ScriptedModel(plan(), PlanReview(outcome="unsupported", questions=[], reason="Individual students are excluded"))
    search = Search()
    report = discover_verified("Find student coordinators' personal contacts", model, search, Fetch())
    assert report.status == "unsupported" and not search.queries


def test_cross_page_institution_evidence_has_source_url():
    p = plan(criteria=[dict(id="city", requirement="In Mumbai", prompt_quote="Mumbai", applies_to="institution")])
    d = decision(evidence=[dict(criterion_id="city", quote="Mumbai campus address")])
    evidence = _check_page(d, p, "Example Engineering College Official faculty directory in Mumbai",
                          [{"url": "https://college.example/contact", "text": "Mumbai campus address"}], [])
    assert evidence[0]["source_url"].endswith("/contact")
    with pytest.raises(ValueError, match="Current page"):
        _check_page(d, p, "A news article", [{"url": URL, "text": HTML}], [])


def test_person_requirement_is_deferred_until_extraction():
    p = plan(criteria=[dict(id="topic", requirement="Researches robotics", prompt_quote="robotics", applies_to="person")], requested_institutions=1)
    report = discover_verified("Find robotics faculty", ScriptedModel(p, REVIEW, decision()), Search(), Fetch(), Limits(rounds=1))
    assert report.candidates  # Directory need not include every faculty biography.


def test_one_site_cannot_consume_the_entire_discovery_budget():
    class Model:
        def ask(self, schema, instruction, payload):
            if schema is Plan:
                return plan(requested_institutions=2)
            if schema is PlanReview:
                return REVIEW
            return decision(accepted="second.example" in payload["url"])
    urls = [f"https://first.example/page-{i}" for i in range(20)] + ["https://second.example/faculty"]
    report = discover_verified("Find two colleges", Model(), Search(urls), Fetch(), Limits(pages=20, institutions=2, rounds=1))
    assert any("second.example" in c["url"] for c in report.candidates)
    assert report.pages <= 20


def test_empty_query_with_partial_engine_outage_does_not_hide_other_results():
    from campusscraper.discovery_services import IncompleteSearch
    class PartialSearch:
        def search(self, query):
            if query == "first":
                raise IncompleteSearch("Some engines unavailable")
            return [URL]
    model = ScriptedModel(plan(queries=["first", "second"], requested_institutions=1), REVIEW, decision())
    report = discover_verified("Find faculty", model, PartialSearch(), Fetch(), Limits(rounds=1))
    assert report.candidates and any("Search incomplete" in note for note in report.issues)


def test_source_template_is_compatible_with_standard_upload(tmp_path):
    from campusscraper.sources import read_sources
    report = Report(prompt="Find faculty", plan=plan().model_dump(), candidates=[{"url": URL, "institution": "Engineering College"}])
    job = {"body": {"prompt": report.prompt, "pages_per_college": 100}, "report": asdict(report)}
    path = export_sources(job, tmp_path / "sources.xlsx")
    sources, errors = read_sources(path)
    assert len(sources) == 1 and sources[0].urls == [URL] and not errors


def person(name, research):
    rec = PersonRecord(URL, relevance_score=1)
    rec.set("name", name, "h1")
    rec.set("research_areas", research, "sections")
    rec.set("email", name.lower().replace(" ", ".") + "@college.example", "labels")
    return rec


def topic_report():
    p = plan(criteria=[dict(id="topic", requirement="Researches robotics", prompt_quote="robotics", applies_to="person")])
    return Report(prompt="Find robotics faculty", plan=p.model_dump(), status="candidates_found")


def test_filter_is_per_person_and_cannot_create_contacts():
    records = [person("Alice Rao", ["Robotics and automation"]), person("Brian Shah", ["Cloud computing"])]
    response = RecordBatchDecision(decisions=[
        {"id": 0, "accepted": True, "reason": "Published research", "evidence": [{"criterion_id": "topic", "quote": "Robotics and automation"}]},
        {"id": 1, "accepted": True, "reason": "Fabricated evidence", "evidence": [{"criterion_id": "topic", "quote": "Robotics and automation"}]},
    ])
    filtered, audit = filter_records(topic_report(), records, ScriptedModel(response))
    assert filtered == [records[0]]
    assert filtered[0].get("email") == records[0].get("email")
    assert not audit[1]["accepted"]


def test_filter_requires_exactly_one_vote_per_record():
    from campusscraper.discovery_services import ConfigurationError
    with pytest.raises(ConfigurationError, match="missing or duplicate"):
        filter_records(topic_report(), [person("Alice Rao", ["Robotics"])], ScriptedModel(RecordBatchDecision(decisions=[])))


def test_source_template_disables_rows_when_person_filters_needed(tmp_path):
    report = topic_report()
    report.candidates = [{"institution": "Example Engineering College", "url": URL}]
    job = {"body": {"prompt": report.prompt, "pages_per_college": 150}, "report": asdict(report)}
    path = export_sources(job, tmp_path / "sources.xlsx")
    wb = load_workbook(path, data_only=True)
    header, row = list(wb["Colleges"].values)
    values = dict(zip(header, row))
    assert values["Enabled"] == "no" and "prompt run" in values["Notes"]
    assert values["Faculty Page URLs"] == URL
    wb.close()


@pytest.fixture
def workflow(tmp_path, monkeypatch, site_url):
    import campusscraper.prompt_jobs as jobs
    monkeypatch.setenv("MODEL_PROVIDERS", "groq,gemini")
    monkeypatch.setenv("GROQ_API_KEY", "fixture-key-never-real")
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.setenv("GROQ_TPM", "100000")  # tests must not wait for real-minute pacing
    monkeypatch.setenv("CAMPUSSCRAPER_ALLOW_PRIVATE", "1")
    monkeypatch.setattr(jobs, "ScrapeConfig", lambda **kw: ScrapeConfig(**{**kw, "delay_seconds": 0, "retries": 0}))
    requests = []
    state = {"fail_filter_once": False, "failed": False}
    class SearchHandler(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass
        def do_GET(self):
            requests.append(self.path)
            data = json.dumps({"results": [{"url": site_url + "/faculty"}]}).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(data)
    server = ThreadingHTTPServer(("127.0.0.1", 0), SearchHandler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    monkeypatch.setenv("SEARXNG_URL", f"http://127.0.0.1:{server.server_address[1]}")
    model_calls = []
    def model_http(req):
        body = json.loads(req.content)
        name = body["format"]["title"] if "format" in body else body["response_format"]["json_schema"]["name"]
        payload = json.loads(body["messages"][1]["content"])
        model_calls.append(name)
        if name == "Plan":
            if payload.get("previous_plan") and state.get("fail_repair_once") and not state["failed"]:
                state["failed"] = True
                return httpx.Response(429, headers={"retry-after": "120"})
            data = plan(requested_institutions=1, criteria=[dict(id="topic", requirement="Researches machine learning", prompt_quote="machine learning", applies_to="person")]).model_dump()
            if state.get("plan_quotes"):
                data["criteria"][0]["prompt_quote"] = state["plan_quotes"].pop(0)
        elif name == "PlanReview":
            data = REVIEW.model_dump()
        elif name == "Decision":
            if state.get("fail_discovery_once") and not state["failed"]:
                state["failed"] = True
                return httpx.Response(429, headers={"retry-after": "120"})
            data = decision(institution="Sunrise Institute of Technology", official_site_quote="Sunrise Institute of Technology", staff_page_quote="Our Faculty").model_dump()
        elif name == "RecordBatchDecision":
            if state["fail_filter_once"] and not state["failed"]:
                state["failed"] = True
                return httpx.Response(429, headers={"retry-after": "120"})
            data = {"decisions": [{"id": r["id"], "accepted": "Anitha" in r["record_name"], "reason": "Published research" if "Anitha" in r["record_name"] else "Topic not supported",
                "evidence": [{"criterion_id": "topic", "quote": "Machine Learning"}] if "Anitha" in r["record_name"] else []} for r in payload["records"]]}
        else:
            raise AssertionError(name)
        if "format" in body:
            return httpx.Response(200, json={"done": True, "done_reason": "stop", "message": {"content": json.dumps(data)}})
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(data)}}], "usage": {"total_tokens": 100}})
    client = httpx.Client(transport=httpx.MockTransport(model_http))
    monkeypatch.setattr(jobs, "ModelRouter", lambda db, job, control: ModelRouter(db, job, control, client=client))
    history = Store(tmp_path / "history.db")
    db = DiscoveryDB(tmp_path / "prompt-discovery.db")
    manager = PromptManager(db, history, tmp_path / "results")
    monkeypatch.setattr(manager, "start_worker", lambda: None)
    yield manager, requests, model_calls, state
    manager.close()
    server.shutdown()
    client.close()
    db.db.close()
    history._db.close()


def test_complete_prompt_to_excel_and_history(workflow):
    manager, searches, calls, _ = workflow
    job = manager.create(PromptIn(engine="evidence", prompt="Find one college and faculty researching machine learning", render="never", pages_per_college=60), start=False)
    manager.execute(job["id"])
    saved = manager.db.get(job["id"])
    assert saved["state"] in {"done", "partial"}, saved
    assert saved["phase"] == "finished"
    wb = load_workbook(saved["xlsx"], read_only=True, data_only=True)
    data = json.dumps(list(wb["Faculty"].values))
    assert "anitha.rao@sunrise-it.ac.in" in data
    assert "Dr. Anitha Rao" in data
    assert "srinivas.k@" not in data and "rahul.v21" not in data
    assert len(manager.history.run_people(job["id"])) == 1
    assert searches and calls[:3] == ["Plan", "PlanReview", "Decision"]
    assert saved["stats"]["filtered_out"] > 0
    wb.close()


def test_repaired_prompt_completes_real_crawl_filtering_and_excel(workflow):
    manager, searches, calls, state = workflow
    state["plan_quotes"] = ["research in machine learning", "machine learning"]
    job = manager.create(PromptIn(engine="evidence", prompt="Find faculty researching machine learning", render="never", pages_per_college=60), start=False)
    manager.execute(job["id"])
    saved = manager.db.get(job["id"])
    assert saved["state"] in {"done", "partial"}, saved
    assert calls[:4] == ["Plan", "Plan", "PlanReview", "Decision"]
    assert len(saved["report"]["planning_attempts"]) == 2 and searches
    wb = load_workbook(saved["xlsx"], read_only=True, data_only=True)
    data = json.dumps(list(wb["Faculty"].values))
    assert "anitha.rao@sunrise-it.ac.in" in data
    assert "srinivas.k@" not in data and "rahul.v21" not in data
    wb.close()


def test_failed_planning_export_includes_failure_and_rejected_criteria(workflow, monkeypatch):
    import campusscraper.web as web
    manager, searches, calls, state = workflow
    state["plan_quotes"] = ["unsupported model paraphrase", "unsupported model paraphrase"]
    monkeypatch.setattr(web, "_PROMPT_MANAGER", manager)
    job = manager.create(PromptIn(engine="evidence", prompt="Find machine learning faculty", render="never"), start=False)
    manager.execute(job["id"])
    client = TestClient(web.app)
    exported = client.get(f"/api/discovery/{job['id']}/evidence.json")
    data = exported.json()
    assert data["job"]["state"] == "failed"
    assert data["job"]["phase"] == data["report"]["status"] == "planning_failed"
    assert "automatic repair" in data["job"]["error"]
    assert data["report"]["model_calls"] == 2
    assert len(data["report"]["planning_attempts"]) == 2
    assert data["report"]["planning_attempts"][0]["grounding_errors"]
    assert "fixture-key" not in exported.text
    assert not searches and calls == ["Plan", "Plan"]
    assert client.get(f"/api/discovery/{job['id']}/download.xlsx").status_code == 409


def test_quota_during_plan_repair_resumes_without_repeating_initial_call(workflow):
    manager, searches, calls, state = workflow
    state.update(plan_quotes=["research in machine learning", "machine learning"], fail_repair_once=True)
    job = manager.create(PromptIn(engine="evidence", prompt="Find machine learning faculty", render="never", pages_per_college=60), start=False)
    manager.execute(job["id"])
    waiting = manager.db.get(job["id"])
    assert waiting["state"] == "waiting" and not searches
    assert waiting["report"]["planning_attempts"][0]["grounding_errors"]
    manager.db.cooldown("groq", 0)
    manager.action(job["id"], "resume")
    manager.execute(job["id"])
    saved = manager.db.get(job["id"])
    assert saved["state"] in {"done", "partial"}, saved
    assert calls.count("Plan") == 3  # Initial + quota response + successful repair.
    assert len(saved["report"]["planning_attempts"]) == 2


def test_discovery_quota_then_process_restart_preserves_page_and_finishes_excel(workflow, monkeypatch):
    import campusscraper.prompt_jobs as jobs
    manager, searches, calls, state = workflow
    state["fail_discovery_once"] = True
    job = manager.create(PromptIn(engine="evidence", prompt="Find machine learning faculty", render="never", pages_per_college=60), start=False)
    manager.execute(job["id"])
    waiting = manager.db.get(job["id"])
    assert waiting["state"] == "waiting"
    assert waiting["report"]["pages"] == 1 and waiting["report"]["status"] == "checking_pages"
    assert waiting["report"]["search_attempts"]
    assert manager.db.cached(job["id"], "discovery_state_v1")["work"]["pending"]["stage"] == "assess"
    before_searches = len(searches)
    # Reopen the actual SQLite journal to simulate a new process, not in-memory resume.
    new_db = DiscoveryDB(manager.db.path)
    resumed = PromptManager(new_db, manager.history, manager.out_dir)
    monkeypatch.setattr(resumed, "start_worker", lambda: None)
    monkeypatch.setattr(jobs.CachedSearch, "search", lambda *a: pytest.fail("Completed search was replayed"))
    monkeypatch.setattr(jobs.CachedFetcher, "get", lambda *a: pytest.fail("Pending page bytes were refetched"))
    reports = []
    original_update = resumed.update
    def capture(job_id, **values):
        if "report" in values:
            reports.append(values["report"])
        return original_update(job_id, **values)
    monkeypatch.setattr(resumed, "update", capture)
    try:
        new_db.cooldown("groq", 0)
        resumed.action(job["id"], "resume")
        resumed.execute(job["id"])
        result = new_db.get(job["id"])
        assert result["state"] in {"done", "partial"}, result
        assert all(r["pages"] >= 1 for r in reports)
        assert len(searches) == before_searches
        assert calls.count("Plan") == calls.count("PlanReview") == 1
        assert calls.count("Decision") == 2
        wb = load_workbook(result["xlsx"], read_only=True, data_only=True)
        data = json.dumps(list(wb["Faculty"].values))
        assert "anitha.rao@sunrise-it.ac.in" in data and "rahul.v21" not in data
        wb.close()
    finally:
        resumed.close()
        new_db.db.close()


def test_unfiltered_prompt_preserves_manual_crawler_contacts(workflow, site_url, tmp_path):
    from campusscraper.batch import run_batch
    from campusscraper.sources import Source
    manager, *_ = workflow
    job = manager.create(PromptIn(engine="evidence", prompt="Find Sunrise faculty", render="never", pages_per_college=60), start=False)
    report = Report(prompt=job["body"]["prompt"], plan=plan(requested_institutions=1).model_dump(),
                    status="candidates_found", candidates=[{"institution": "Sunrise Institute of Technology", "url": site_url + "/faculty"}])
    manager.db.cache(job["id"], "discovery_report", asdict(report))
    manager.execute(job["id"])
    assert manager.db.get(job["id"])["state"] in {"done", "partial"}
    live = {}
    run_batch([Source(2, "Sunrise Institute of Technology", [site_url + "/faculty"])], tmp_path / "manual.xlsx",
              ScrapeConfig(render="never", allow_private_hosts=True, delay_seconds=0, retries=0, workers=2),
              60, live=live, resume=False, parallel_sites=2)
    fields = ("name", "email", "phone", "contact_email", "contact_phone")
    expected = {tuple(d.get(f) for f in fields) for d in live["merger"].records if d.is_relevant}
    actual = {tuple(d.get(f) for f in fields) for d in manager.history.run_people(job["id"])}
    assert actual == expected and len(actual) >= 8


def test_quota_during_filtering_restarts_without_search_or_crawl(workflow, monkeypatch):
    import campusscraper.prompt_jobs as jobs
    manager, searches, calls, state = workflow
    state["fail_filter_once"] = True
    job = manager.create(PromptIn(engine="evidence", prompt="Find machine learning faculty", render="never", pages_per_college=60), start=False)
    manager.execute(job["id"])
    assert manager.db.get(job["id"])["state"] == "waiting"
    assert not manager.history.run_people(job["id"])
    assert manager.db.cached(job["id"], "extracted")
    before = len(searches)
    monkeypatch.setattr(jobs, "run_batch", lambda *a, **kw: pytest.fail("Finished extraction was repeated"))
    manager.db.cooldown("groq", 0)
    manager.action(job["id"], "resume")
    manager.execute(job["id"])
    assert manager.db.get(job["id"])["state"] in {"done", "partial"}
    assert len(searches) == before
    assert calls.count("Plan") == 1 and calls.count("Decision") == 1


def test_pause_prevents_work_and_can_resume(workflow):
    manager, searches, _, _ = workflow
    job = manager.create(PromptIn(engine="evidence", prompt="Find machine learning faculty", render="never"), start=False)
    assert manager.action(job["id"], "pause")["state"] == "paused"
    assert not searches
    manager.action(job["id"], "resume")
    assert manager.db.get(job["id"])["state"] == "queued"


def test_restart_marks_active_job_interrupted(workflow):
    manager, *_ = workflow
    job = manager.create(PromptIn(engine="evidence", prompt="Find machine learning faculty"), start=False)
    manager.update(job["id"], state="running")
    new = PromptManager(manager.db, manager.history, manager.out_dir)
    assert new.db.get(job["id"])["state"] == "interrupted"


def test_portal_api_hides_unfiltered_results_and_keys(workflow, monkeypatch):
    import campusscraper.web as web
    manager, *_ = workflow
    monkeypatch.setattr(web, "_STORE", manager.history)
    monkeypatch.setattr(web, "_PROMPT_MANAGER", manager)
    monkeypatch.setattr(web, "OUT_DIR", manager.out_dir)
    client = TestClient(web.app)
    config = client.get("/api/discovery/config")
    assert config.status_code == 200 and "fixture-key" not in config.text
    response = client.post("/api/discovery", json={"prompt": "Find machine learning faculty", "render": "never", "pages_per_college": 60})
    assert response.status_code == 200
    job_id = response.json()["id"]
    assert client.get(f"/api/discovery/{job_id}/download.xlsx").status_code == 409
    manager.execute(job_id)
    status = client.get(f"/api/discovery/{job_id}").json()
    assert status["download_ready"] and "xlsx" not in status
    assert client.get(f"/api/runs/{job_id}/download.xlsx").status_code == 200
    assert client.get(f"/api/discovery/{job_id}/sources.xlsx").status_code == 200
    assert client.get(f"/api/discovery/{job_id}/evidence.json").status_code == 200
    assert client.post("/api/discovery", json={"prompt": "   "}).status_code == 422
    assert client.post("/api/discovery", json={"prompt": "faculty", "search_url": "http://evil"}).status_code == 422


def test_ollama_prompt_to_real_crawler_and_excel(workflow, monkeypatch):
    # Real crawler/export with mocked model transport: no cloud API required.
    monkeypatch.setenv("MODEL_PROVIDERS", "ollama")
    test_complete_prompt_to_excel_and_history(workflow)


def test_sites_engine_prompt_to_excel_without_calling_any_model(workflow):
    """The whole chain - prompt, search, verified URLs, contacts, Excel.

    A model may be asked to read the request, once. Nothing else in this engine
    asks a model anything: the sites themselves decide which colleges qualify.
    """
    manager, searches, calls, _ = workflow
    job = manager.create(PromptIn(prompt="Find 1 engineering college in Hyderabad and return its faculty page URLs",
                                  render="never", pages_per_college=60, site_pages=10), start=False)
    manager.execute(job["id"])
    saved = manager.db.get(job["id"])
    # "partial" is the honest state when a site reports dead links; both carry results.
    assert saved["state"] in {"done", "partial"}, saved.get("error")
    assert calls in ([], ["RequestRead"]), f"a model judged something: {calls}"
    report = saved["report"]
    assert report["status"] == "candidates_found"
    assert report["plan"]["request"]["kind"] == "engineering"
    assert report["plan"]["request"]["location"] == "Hyderabad"
    college = report["plan"]["colleges"][0]
    assert "Sunrise Institute of Technology" in college["name"]
    assert college["city"] == "Hyderabad"
    assert college["staff_found"] >= 5
    assert any(url.endswith("/departments/cse/faculty") or "faculty" in url
               for url in college["faculty_urls"])
    candidate = report["candidates"][0]
    assert candidate["institution"] == college["name"] and candidate["extracted_staff"] >= 5
    workbook = load_workbook(saved["xlsx"], read_only=True, data_only=True)
    names = {row[0].value for row in workbook["Faculty"].iter_rows(min_row=2)}
    assert "Dr. Anitha Rao" in names


def test_sites_engine_states_what_it_did_not_filter(workflow):
    manager, _, calls, _ = workflow
    job = manager.create(PromptIn(prompt="Find 2 engineering colleges in Hyderabad with faculty researching machine learning",
                                  render="never", auto_scrape=False, site_pages=8), start=False)
    manager.execute(job["id"])
    report = manager.db.get(job["id"])["report"]
    assert calls in ([], ["RequestRead"]), f"a model judged something: {calls}"
    assert any("not applied" in issue for issue in report["issues"])
    assert any("Requested 2 colleges" in issue for issue in report["issues"])
