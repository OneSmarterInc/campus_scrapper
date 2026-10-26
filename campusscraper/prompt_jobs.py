"""Durable prompt → discovery → existing crawler → filtered Excel workflow."""
from __future__ import annotations

import json
import threading
import time
import uuid
from dataclasses import asdict
from pathlib import Path
from typing import Literal

from pydantic import BaseModel, ConfigDict, Field, field_validator

from .batch import run_batch
from .config import ScrapeConfig
from .crawl import CrawlResult, Crawler, PageLog
from .discovery_services import (
    CachedFetcher, CachedSearch, ConfigurationError, DiscoveryDB, ModelRouter, Paused, ServiceWait, settings,
)
from .export import to_xlsx
from .fetch import Fetcher
from .models import OfficeRecord, PersonRecord
from .prompt_discovery import Limits, Report
from .site_discovery import College, discover_colleges, read_request, sources_from_colleges
from .verified_discovery import discover_verified, filter_records, sources_from_report

ACTIVE = {"queued", "running", "waiting", "pausing", "stopping"}
TERMINAL = {"done", "partial", "clarify", "unsupported", "no_candidates", "sources_ready"}


class PromptIn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    prompt: str = Field(min_length=1, max_length=6000)
    max_colleges: int = Field(10, ge=1, le=100)
    discovery_pages: int = Field(60, ge=1, le=100)
    searches: int = Field(12, ge=1, le=30)
    pages_per_college: int = Field(200, ge=1, le=1000)
    render: Literal["never", "auto", "always"] = "auto"
    auto_scrape: bool = True
    # "sites" proves each college by crawling it; "evidence" asks a model to judge pages.
    engine: Literal["sites", "evidence"] = "sites"
    site_pages: int = Field(14, ge=2, le=60)

    @field_validator("prompt")
    @classmethod
    def nonempty(cls, value):
        if not value.strip():
            raise ValueError("Describe the faculty, colleges or offices you need")
        return value.strip()


class PromptManager:
    def __init__(self, db, history, out_dir):
        self.db, self.history, self.out_dir = db, history, Path(out_dir)
        self.lock = threading.RLock()
        self.wake = threading.Event()
        self.closed = threading.Event()
        self.thread = None
        self.current = None
        self.stop = threading.Event()
        # A new process cannot still own an in-flight HTTP/browser operation.
        for job in self.db.jobs():
            if job["state"] in {"running", "pausing", "stopping"}:
                job["state"] = "interrupted"
                self.db.save(job)
                self.history.update_run(job["id"], state="interrupted")

    def start_worker(self):
        with self.lock:
            if self.thread is None or not self.thread.is_alive():
                self.thread = threading.Thread(target=self._loop, daemon=True, name="prompt-discovery")
                self.thread.start()
        self.wake.set()

    def close(self):
        self.closed.set()
        self.stop.set()
        self.wake.set()

    def create(self, body, start=True):
        config = settings()
        if not (config["ollama_configured"] or config["groq_configured"] or config["gemini_configured"]):
            raise ConfigurationError("Configure Ollama or an explicitly selected cloud provider in .env and restart")
        job = {"id": uuid.uuid4().hex[:10], "state": "queued", "phase": "queued", "started": time.time(),
               "finished": None, "body": body.model_dump(), "report": {}, "rows": [], "progress": {},
               "error": None, "retry_at": None, "xlsx": None, "stats": {}}
        self.history.create_run(job["id"], "prompt", body.prompt[:120], body.model_dump())
        self.history.update_run(job["id"], state="queued")
        self.db.save(job)
        if start:
            self.start_worker()
        return job

    def update(self, job_id, **values):
        with self.lock:
            job = self.db.get(job_id)
            if job is None:
                return
            job.update(values)
            self.db.save(job)
            hist = {k: values[k] for k in ("state", "finished", "error", "xlsx", "stats") if k in values}
            if hist:
                self.history.update_run(job_id, **hist)

    def action(self, job_id, action):
        with self.lock:
            job = self.db.get(job_id)
            if not job:
                raise KeyError("Prompt run not found")
            state = job["state"]
            if action in {"pause", "stop"}:
                if state not in ACTIVE:
                    raise ValueError(f"Cannot {action} a run that is {state}")
                active = self.current == job_id
                next_state = ("pausing" if action == "pause" else "stopping") if active else ("paused" if action == "pause" else "stopped")
                self.update(job_id, state=next_state, retry_at=None)
                if active:
                    self.stop.set()
            elif action == "resume":
                if state not in {"paused", "stopped", "interrupted", "waiting", "failed"}:
                    raise ValueError(f"Cannot resume a run that is {state}")
                self.update(job_id, state="queued", error=None, retry_at=None, finished=None)
                self.start_worker()
            else:
                raise ValueError("Unknown action")
        self.wake.set()
        return self.db.get(job_id)

    def _loop(self):
        while not self.closed.is_set():
            with self.lock:
                pending = [j for j in reversed(self.db.jobs()) if j["state"] == "queued" or
                           (j["state"] == "waiting" and (j.get("retry_at") or 0) <= time.time())]
                job = pending[0] if pending else None
                if job:
                    self.current = job["id"]
                    self.stop = threading.Event()
                    self.update(job["id"], state="running", error=None, retry_at=None)
            if job:
                self.execute(job["id"])
                with self.lock:
                    self.current = None
            else:
                self.wake.wait(2)
                self.wake.clear()

    def _discover_sites(self, body, search, progress, control):
        """Find colleges and their faculty pages by crawling each site.

        Search only suggests where to look. A college appears in the result
        because our extractor read its staff from one of its own pages.
        """
        request = read_request(body.prompt)
        want = min(request.count, body.max_colleges)
        report = Report(prompt=body.prompt)
        criteria = [{"id": "institution_type", "requirement": f"{request.kind} college",
                     "prompt_quote": request.kind, "applies_to": "institution"}] if request.kind else []
        if request.location:
            criteria.append({"id": "location", "requirement": f"Located in {request.location}",
                             "prompt_quote": request.location, "applies_to": "institution"})
        report.plan = {"objective": body.prompt, "target": request.target, "criteria": criteria,
                       "requested_institutions": request.count, "explanation": request.summary(),
                       "request": request.as_dict(), "colleges": []}
        if request.count > body.max_colleges:
            report.issues.append(f"Requested {request.count} colleges but this run is limited to "
                                 f"{body.max_colleges}. Raise 'Maximum colleges' to {request.count}")
        if request.departments:
            report.issues.append("Keeping only these departments in the contact export: "
                                 + ", ".join(request.departments))
        if request.unfiltered_conditions:
            report.issues.append("This run verifies colleges and faculty pages. Conditions about individual "
                                 "people (research area, email required, experience) were not applied; "
                                 "switch the engine to 'AI judges pages' for those")

        def make_crawler():
            return Crawler(ScrapeConfig(render=body.render, strict_public_network=True, max_bytes=2_000_000,
                                        retries=1, max_pages=body.site_pages, workers=2, patience=0,
                                        include_offices=request.target != "faculty"),
                           stop_event=self.stop)

        colleges, rejected, notes = discover_colleges(
            search, make_crawler, want=want, kind=request.kind, location=request.location,
            target=request.target, names=request.names, max_searches=body.searches,
            control=control, progress=progress)
        report.plan["colleges"] = [c.as_dict() for c in colleges]
        report.searches = len(getattr(search, "queries", []))
        report.pages = sum(c.pages_read for c in colleges)
        for college in colleges:
            urls = college.office_urls if request.target == "offices" else college.faculty_urls
            for url in urls:
                report.candidates.append({
                    "url": url, "institution": college.name, "accepted": True,
                    "extracted_staff": college.staff_found,
                    "reason": f"{college.staff_found} staff read from this site by the extractor",
                    "evidence_sources": [{"criterion_id": k, "quote": v, "source_url": college.homepage}
                                         for k, v in college.evidence.items()]})
        report.attempts = [{"url": f"https://{r.host}/", "institution": r.name, "accepted": False,
                            "reason": r.reason, "evidence_sources": []} for r in rejected]
        report.issues.extend(notes)
        report.issues = list(dict.fromkeys(report.issues))
        report.status = "candidates_found" if report.candidates else "no_candidates"
        return report

    def execute(self, job_id):
        """Public for deterministic integration tests; production uses one worker."""
        job = self.db.get(job_id)
        body = PromptIn(**job["body"])
        model = search = fetcher = None

        def control():
            if self.closed.is_set() or self.stop.is_set():
                raise Paused()

        def progress(event):
            phase = event.pop("phase", "checking_pages")
            changes = {"phase": phase, "progress": event}
            if "report" in event:
                changes["report"] = event.pop("report")
            self.update(job_id, **changes)
            control()  # Persist the operation that just finished even if Pause was clicked.

        def fail(message):
            latest = self.db.get(job_id)
            report = latest.get("report") or asdict(Report(prompt=body.prompt))
            planning = latest.get("phase") in {"planning", "planning_failed"}
            report["status"] = "planning_failed" if planning else "failed"
            if message not in report.setdefault("issues", []):
                report["issues"].append(message)
            self.update(job_id, state="failed", phase="planning_failed" if planning else latest["phase"],
                        report=report, error=message, finished=time.time())

        try:
            self.update(job_id, state="running")
            model = ModelRouter(self.db, job_id, control)
            saved = self.db.cached(job_id, "discovery_report")
            if saved:
                report = Report(**saved)
            else:
                config = settings()
                search = CachedSearch(self.db, job_id, config["search_url"], control)
                fetcher = CachedFetcher(Fetcher(ScrapeConfig(render=body.render, strict_public_network=True,
                    max_bytes=2_000_000, retries=1), self.stop), self.db, job_id, control)
                if body.engine == "sites":
                    report = self._discover_sites(body, search, progress, control)
                else:
                    report = discover_verified(body.prompt, model, search, fetcher,
                        Limits(searches=body.searches, pages=body.discovery_pages, institutions=body.max_colleges, rounds=2), control, progress,
                        load_state=lambda: self.db.cached(job_id, "discovery_state_v1"),
                        save_state=lambda state: self.db.cache(job_id, "discovery_state_v1", state))
                self.db.cache(job_id, "discovery_report", asdict(report))
            self.update(job_id, report=asdict(report))
            if report.status != "candidates_found":
                self.update(job_id, state=report.status, phase=report.status, finished=time.time())
                return
            sources = (sources_from_colleges([College(**c) for c in report.plan["colleges"]],
                                             report.plan.get("target", "faculty"))
                       if body.engine == "sites" else sources_from_report(report))
            if not body.auto_scrape:
                self.update(job_id, state="sources_ready", phase="sources_ready", finished=time.time())
                return
            extracted = self.db.cached(job_id, "extracted")
            if extracted is None:
                rows = {s.row: {"sequence": i, "row": s.row, "institution": s.institution,
                               "status": "queued", "pages": 0, "faculty": 0} for i, s in enumerate(sources, 1)}
                self.update(job_id, phase="scraping", rows=list(rows.values()))
                live, events_lock = {}, threading.RLock()

                def crawl_event(e):
                    with events_lock:
                        row = rows.get(e.get("row"))
                        if row is not None:
                            if e["event"] == "site_start":
                                row["status"] = "scraping"
                            elif e["event"] == "page":
                                row.update(pages=e["fetched"], faculty=e.get("faculty", 0))
                            elif e["event"] == "site_done":
                                row.update({k: v for k, v in e.items() if k != "event"})
                        self.update(job_id, rows=list(rows.values()), progress={"completed": sum(r["status"] in {"done", "resumed", "failed", "error"} for r in rows.values()), "total": len(rows)})

                self.out_dir.mkdir(parents=True, exist_ok=True)
                run_batch(sources, self.out_dir / f"prompt-{job_id}-raw.xlsx",
                    ScrapeConfig(render=body.render, strict_public_network=True, workers=2,
                                 include_offices=True,
                                 departments=list((report.plan.get("request") or {}).get("departments") or [])),
                    body.pages_per_college, resume=True, progress=crawl_event, stop_event=self.stop,
                    live=live, parallel_sites=2)
                control()
                if len(live["summary"]) != len(sources):
                    raise RuntimeError("Some college workers did not finish. Resume to retry")
                extracted = {"doctors": [d.to_full_dict() for d in live["merger"].records],
                             "offices": [o.to_full_dict() for o in live.get("offices", [])],
                             "pages": [asdict(p) for p in live["pages"]], "sites": live["summary"]}
                self.db.cache(job_id, "extracted", extracted)
            control()
            target = report.plan.get("target", "faculty")
            docs = [PersonRecord.from_full_dict(d) for d in extracted["doctors"]]
            docs = [] if target == "offices" else (docs if target == "staff" else [d for d in docs if d.is_relevant])
            offices = [OfficeRecord.from_full_dict(o) for o in extracted["offices"]] if target in {"offices", "both"} else []
            self.update(job_id, phase="filtering")
            filtered, audit = filter_records(report, [*docs, *offices], model, control, progress)
            docs = [d for d in filtered if isinstance(d, PersonRecord)]
            offices = [o for o in filtered if isinstance(o, OfficeRecord)]
            docs.sort(key=lambda d: (d.source_institution.casefold(), (d.get("name") or "").casefold()))
            per_college = report.plan.get("requested_people_per_institution")
            if per_college:
                selected = []
                for source in sources:
                    matched = [d for d in docs if d.source_institution == source.institution]
                    if len(matched) < per_college:
                        report.issues.append(f"{source.institution}: requested {per_college} people; {len(matched)} matched")
                    selected.extend(matched[:per_college])
                docs = selected
            count = report.plan.get("requested_people")
            if count and len(docs) < count:
                report.issues.append(f"Requested {count} people; {len(docs)} matched with available evidence")
            if count:
                docs = docs[:count]
            if any(a.get("evidence_window_limited") for a in audit):
                report.issues.append("Some long record descriptions exceeded the filtering evidence window; coverage may be limited")
            matching_colleges = {d.source_institution for d in docs} | {o.source_institution for o in offices}
            requested_colleges = report.plan.get("requested_institutions")
            if requested_colleges and len(matching_colleges) < requested_colleges:
                report.issues.append(f"Requested {requested_colleges} colleges; {len(matching_colleges)} produced matching records after extraction and filtering")
            if not docs and not offices:
                report.issues.append("No matching records were extracted; source discovery alone does not establish contact coverage")
            sites = []
            for site in extracted["sites"]:
                matched = [d for d in docs if d.source_institution == site["institution"]]
                site = {**site, "people": len(matched), "faculty": sum(d.is_relevant for d in matched),
                        "offices": sum(o.source_institution == site["institution"] for o in offices),
                        "with_email": sum(bool(d.get("email")) for d in matched),
                        "with_phone": sum(bool(d.get("phone")) for d in matched),
                        "with_any_contact": sum(any(d.get(f) for f in ("email", "phone", "contact_email", "contact_phone")) for d in matched)}
                sites.append(site)
            limited = any(s["status"] in {"failed", "error"} or s.get("note") for s in sites) or bool(report.issues)
            result = CrawlResult([u for s in sources for u in s.urls], docs, [PageLog(**p) for p in extracted["pages"]], offices=offices)
            result.stats = {"prompt": body.prompt, "colleges": len(sources), "people_found": len(docs), "faculty": sum(d.is_relevant for d in docs), "student_offices": len(offices),
                            "colleges_with_matching_records": len(matching_colleges),
                            "pages_fetched": len(result.pages), "filtered_out": sum(not a["accepted"] for a in audit),
                            "coverage": "partial / limited" if limited else "configured run completed; not a guarantee of exhaustive coverage",
                            "discovery_notes": " | ".join(report.issues), "seconds": round(time.time() - job["started"], 1)}
            control()
            self.update(job_id, phase="exporting", progress={})
            path = to_xlsx(result, self.out_dir / f"prompt-{job_id}.xlsx", sites=sites)
            self.db.cache(job_id, "filter_audit", audit)
            self.history.save_people(job_id, docs)
            self.history.save_offices(job_id, offices)
            for site in sites:
                self.history.save_site(job_id, site)
            self.update(job_id, state="partial" if limited else "done", phase="finished", xlsx=str(path),
                        report=asdict(report), stats=result.stats, finished=time.time())
        except ServiceWait as exc:
            if self.stop.is_set() or self.closed.is_set():
                state = self.db.get(job_id)["state"]
                self.update(job_id, state="stopped" if state == "stopping" else ("interrupted" if self.closed.is_set() else "paused"), retry_at=None)
            else:
                self.update(job_id, state="waiting", error=str(exc), retry_at=exc.retry_at)
        except Paused:
            state = self.db.get(job_id)["state"]
            self.update(job_id, state="stopped" if state == "stopping" else ("interrupted" if self.closed.is_set() else "paused"), retry_at=None)
        except ConfigurationError as exc:
            fail(str(exc))
        except Exception as exc:
            # Do not expose arbitrary HTTP exception bodies or API credentials.
            fail(f"{type(exc).__name__}: the run could not finish. Progress is saved; resume after checking setup")
        finally:
            self.update(job_id, usage=self.db.stats(job_id))
            for service in (model, search, fetcher):
                if service:
                    service.close()


def export_sources(job, path):
    """The same upload columns; filtered prompts cannot silently broaden on re-upload."""
    from openpyxl import Workbook
    from .export import _csv_safe
    from .sources import TEMPLATE_COLUMNS
    report = Report(**job["report"])
    sources = sources_from_report(report)
    requires_filter = any(c["applies_to"] != "institution" for c in report.plan.get("criteria", []))
    wb = Workbook()
    ws = wb.active
    ws.title = "Colleges"
    ws.append([c[1] for c in TEMPLATE_COLUMNS])
    for source in sources:
        url_column = "student_office_urls" if report.plan.get("target") == "offices" else "faculty_page_urls"
        row = {"college_name": source.institution, url_column: "\n".join(source.urls),
               "enabled": "no" if requires_filter else "yes", "max_pages": job["body"]["pages_per_college"],
               "notes": ("Use the prompt run to apply individual/page filters; this row is disabled to prevent an unfiltered crawl. " if requires_filter else "") + job["body"]["prompt"]}
        ws.append([_csv_safe(row.get(c[0], "")) for c in TEMPLATE_COLUMNS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    from openpyxl.styles import Font, PatternFill, Alignment
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="123A44")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 36
        for cell in col:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return Path(path)
