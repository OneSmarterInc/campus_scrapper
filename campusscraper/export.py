"""Excel / CSV / JSON output."""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .crawl import CrawlResult
from .models import ALL_FIELDS, LIST_FIELDS, PersonRecord

COLUMNS = ["name", "designation", "department", "specialty", "institution", "city",
           "email", "phone", "contact_email", "contact_phone",
           "qualifications", "experience_years", "research_areas", "subjects",
           "google_scholar", "orcid", "scopus_id", "researchgate", "linkedin", "vidwan", "homepage",
           "education", "awards", "memberships", "about", "photo_url"]
HEADERS = {"phone": "Faculty Phone", "email": "Faculty Email", "contact_phone": "Department / Office Phone",
           "contact_email": "Department / Office Email", "specialty": "Specialization", "orcid": "ORCID",
           "scopus_id": "Scopus ID", "researchgate": "ResearchGate", "linkedin": "LinkedIn", "vidwan": "Vidwan / IRINS",
           "google_scholar": "Google Scholar", "homepage": "Homepage", "research_areas": "Research Areas"}
assert set(COLUMNS) == set(ALL_FIELDS)
META_COLUMNS = ["source_institution", "profile_url", "source_url", "found_on", "is_relevant", "relevance_score",
                "relevance_terms", "page_type", "fields_filled"]
META_HEADERS = {"source_institution": "Source College", "is_relevant": "In Faculty Sheet", "relevance_score": "Score",
                "relevance_terms": "Why"}
OFFICE_COLUMNS = ["office_type", "office_name", "institution", "city", "emails", "phones", "coordinators",
                  "higher_studies_lines", "page_url", "source_institution", "skipped_student_lines"]
OFFICE_HEADERS = {"office_type": "Office Type", "office_name": "Office / Body", "emails": "Office Emails",
                  "phones": "Office Phones", "coordinators": "Faculty / Staff Coordinators",
                  "higher_studies_lines": "Published Higher-Studies Figures", "page_url": "Page URL",
                  "source_institution": "Source College", "skipped_student_lines": "Student Lines Left Out"}


def office_row(o) -> list:
    d = o.to_dict()
    return [_cell(d.get(c)) if c != "skipped_student_lines" else d.get(c, 0) for c in OFFICE_COLUMNS]


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, list):
        value = "; ".join(str(v) for v in value)
    return str(value)[:32000]


def _csv_safe(value):
    value = str(value)
    if value[:1] in ("=", "+", "-", "@"):
        value = "'" + value  # never let page text become a spreadsheet formula in CSV
    return value


def row(rec: PersonRecord) -> list:
    d = rec.to_dict()
    return [_cell(d.get(c)) for c in COLUMNS] + [
        _cell(rec.source_institution), _cell(rec.profile_url), _cell(rec.source_url), _cell(rec.found_on), "YES" if rec.is_relevant else "no",
        round(rec.relevance_score, 2), _cell(rec.relevance_terms), rec.page_type, len(rec.filled_fields()),
    ]


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"  # page text stays text, never a formula; "+91 ..." shows as is
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="164E63")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns):
        letter = col[0].column_letter
        width = (widths or {}).get(headers[i]) or min(60, max(10, max(len(str(c.value or "")) for c in col[:200]) + 2))
        ws.column_dimensions[letter].width = width
    return ws


def to_xlsx(result: CrawlResult, path: str | Path, sites: list[dict] | None = None) -> Path:
    path = Path(path)
    wb = Workbook()
    wb.remove(wb.active)
    headers = [HEADERS.get(c) or META_HEADERS.get(c) or c.replace("_", " ").title() for c in COLUMNS + META_COLUMNS]
    widths = {"About": 60, "Research Areas": 45, "Profile Url": 45, "Source Url": 45, "Found On": 40}
    onco = sorted(result.faculty, key=lambda d: (d.get("institution") or "", d.get("department") or "", d.get("name") or ""))
    _sheet(wb, "Faculty", headers, [row(d) for d in onco], widths)
    _sheet(wb, "All people", headers, [row(d) for d in sorted(result.doctors, key=lambda d: d.get("name") or "")], widths)
    offices = sorted(getattr(result, "offices", []) or [], key=lambda o: (o.source_institution or o.institution, o.office_type))
    _sheet(wb, "Student offices", [OFFICE_HEADERS.get(c) or c.replace("_", " ").title() for c in OFFICE_COLUMNS],
           [office_row(o) for o in offices],
           {"Office / Body": 36, "Faculty / Staff Coordinators": 50, "Published Higher-Studies Figures": 60, "Page URL": 50})
    ev_rows = []
    # One row per doctor and field: several values (e.g. a list of awards) are joined, not repeated.
    for d in result.doctors:
        grouped: dict[str, dict] = {}
        for e in d.evidence:
            g = grouped.setdefault(e.field, {"values": [], "layers": [], "snippet": e.snippet, "urls": []})
            for bucket, item in (("values", _cell(e.value)), ("layers", e.layer), ("urls", e.source_url)):
                if item and item not in g[bucket]:
                    g[bucket].append(item)
        for field_name, g in grouped.items():
            ev_rows.append([d.get("name"), field_name, "; ".join(map(str, g["values"])), ", ".join(g["layers"]),
                            _cell(g["snippet"]), " ".join(g["urls"])])
    for o in offices:
        for e in o.evidence:
            ev_rows.append([o.office_name, e.field, _cell(e.value), e.layer, _cell(e.snippet), e.source_url])
    _sheet(wb, "Evidence", ["Person / Office", "Field", "Value", "Found by", "Page text", "Page URL"], ev_rows,
           {"Page text": 70, "Value": 40, "Page URL": 45})
    _sheet(wb, "Pages", ["URL", "HTTP status", "Page type", "People on page", "Depth", "Note"],
           [[p.url, p.status, p.page_type, p.doctors, p.depth, p.note] for p in result.pages], {"URL": 70})
    summary = [["Start URL(s)", ", ".join(result.start_urls)], ["Exported at", datetime.now().strftime("%Y-%m-%d %H:%M")]]
    summary += [[k.replace("_", " "), json.dumps(v) if isinstance(v, dict) else v] for k, v in result.stats.items()]
    fill = {f: sum(1 for d in onco if d.get(f)) for f in COLUMNS}
    summary.append(["", ""])
    summary.append(["Field coverage (faculty sheet)", f"{len(onco)} records"])
    summary += [[f, f"{n} ({(100 * n // len(onco)) if onco else 0}%)"] for f, n in fill.items()]
    _sheet(wb, "Summary", ["Item", "Value"], summary, {"Item": 32, "Value": 80})
    wb.move_sheet("Summary", offset=-(len(wb.sheetnames) - 1))
    if sites:
        keys = ["row", "institution", "status", "pages", "people", "faculty", "offices", "with_phone", "with_email",
                "with_any_contact", "start_urls", "note"]
        _sheet(wb, "Colleges", [k.replace("_", " ").title() for k in keys],
               [[_cell(s.get(k)) for k in keys] for s in sites], {"Start Urls": 70, "Note": 50})
        wb.move_sheet("Colleges", offset=-(len(wb.sheetnames) - 2))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def to_csv(doctors: list[PersonRecord]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLUMNS + META_COLUMNS)
    for d in doctors:
        w.writerow([_csv_safe(v) for v in row(d)])
    return buf.getvalue()


def to_json(result: CrawlResult, relevant_only: bool = True) -> str:
    docs = result.faculty if relevant_only else result.doctors
    return json.dumps({"start_urls": result.start_urls, "stats": result.stats,
                       "people": [d.to_dict() for d in docs],
                       "offices": [o.to_dict() for o in getattr(result, "offices", [])]}, ensure_ascii=False, indent=2)


def write(result: CrawlResult, out: str, relevant_only: bool = True) -> Path:
    p = Path(out)
    suffix = p.suffix.lower()
    if suffix == ".xlsx":
        return to_xlsx(result, p)
    p.parent.mkdir(parents=True, exist_ok=True)
    if suffix == ".csv":
        p.write_text(to_csv(result.faculty if relevant_only else result.doctors), encoding="utf-8-sig")
    elif suffix == ".json":
        p.write_text(to_json(result, relevant_only), encoding="utf-8")
    else:
        raise ValueError("output must end in .xlsx, .csv or .json")
    return p


__all__ = ["to_xlsx", "to_csv", "to_json", "write", "COLUMNS", "LIST_FIELDS"]
